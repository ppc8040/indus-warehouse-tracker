from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from werkzeug.utils import secure_filename


DATA_START_ROW = 13
BASE_DATA_END_ROW = 37
BASE_TOTAL_ROW = 38
BASE_PRINT_END_ROW = 68

TEMPLATE_PATH = (
    Path(__file__).resolve().parents[1]
    / "resources"
    / "advance_shipment_template.xlsx"
)


class ShipmentExportError(ValueError):
    pass


def _text(value):
    return "" if value is None else str(value).strip()


def _loading_date(value):
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = _text(value)

    if not text:
        raise ShipmentExportError(
            "The completed loading date is missing."
        )

    try:
        return datetime.fromisoformat(
            text.replace("Z", "+00:00")
        ).date()
    except ValueError as exc:
        raise ShipmentExportError(
            f"Invalid completed loading date: {text}"
        ) from exc


def _positive_integer(value, field_name):
    try:
        number = Decimal(_text(value))
    except (InvalidOperation, ValueError) as exc:
        raise ShipmentExportError(
            f"{field_name} must be a number."
        ) from exc

    if number <= 0 or number != number.to_integral_value():
        raise ShipmentExportError(
            f"{field_name} must be a positive whole number."
        )

    return int(number)


def _part_prefix(part):
    match = re.match(r"^[A-Za-z]{2}", part)

    if not match:
        raise ShipmentExportError(
            f"Cannot determine Description of Goods for part {part}."
        )

    return match.group(0).upper()


def _template_description_map(ws):
    mapping = {}

    for row in range(DATA_START_ROW, BASE_DATA_END_ROW + 1):
        description = _text(ws.cell(row, 2).value)
        part = _text(ws.cell(row, 4).value).upper()

        if not description or not part:
            continue

        prefix = _part_prefix(part)
        previous = mapping.get(prefix)

        if previous and previous != description:
            raise ShipmentExportError(
                "The Excel template contains conflicting descriptions "
                f"for part prefix {prefix}."
            )

        mapping[prefix] = description

    return mapping


def _copy_row_format(ws, source_row, target_row):
    ws.row_dimensions[target_row].height = (
        ws.row_dimensions[source_row].height
    )

    for column in range(1, 15):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)

        if source.has_style:
            target._style = copy(source._style)

        target.number_format = source.number_format
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def _insert_shipment_rows(ws, amount):
    if amount <= 0:
        return

    insert_at = BASE_TOTAL_ROW
    old_max_row = ws.max_row

    merged_ranges = []
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= insert_at:
            merged_ranges.append(str(merged))
            ws.unmerge_cells(str(merged))
        elif merged.max_row >= insert_at:
            raise ShipmentExportError(
                f"Template merge {merged} crosses the insertion row."
            )

    row_settings = {}

    for row in range(insert_at, old_max_row + 1):
        dimension = ws.row_dimensions[row]
        row_settings[row] = {
            "height": dimension.height,
            "hidden": dimension.hidden,
            "outlineLevel": dimension.outlineLevel,
            "collapsed": dimension.collapsed,
        }

    ws.insert_rows(insert_at, amount)

    for old_row, settings in row_settings.items():
        new_dimension = ws.row_dimensions[old_row + amount]

        for name, value in settings.items():
            setattr(new_dimension, name, value)

    for merged in merged_ranges:
        min_col, min_row, max_col, max_row = range_boundaries(
            merged
        )

        ws.merge_cells(
            start_row=min_row + amount,
            start_column=min_col,
            end_row=max_row + amount,
            end_column=max_col,
        )

    for row in range(insert_at, insert_at + amount):
        _copy_row_format(ws, BASE_DATA_END_ROW, row)


def _normalise_items(items, description_map):
    if not isinstance(items, (list, tuple)) or not items:
        raise ShipmentExportError(
            "The loading has no shipment items."
        )

    normalised = []
    seen_hus = set()

    for position, raw in enumerate(items, start=1):
        if not isinstance(raw, dict):
            raise ShipmentExportError(
                f"Loading item {position} is invalid."
            )

        hu = _text(raw.get("handling_unit"))
        part = _text(raw.get("indus_part")).upper()
        customer = _text(raw.get("customer_part"))
        quantity = _positive_integer(
            raw.get("pallet_qty"),
            f"Quantity for item {position}",
        )

        if not hu:
            raise ShipmentExportError(
                f"Handling Unit is missing for item {position}."
            )

        if hu in seen_hus:
            raise ShipmentExportError(
                f"Handling Unit {hu} appears more than once."
            )

        if not part:
            raise ShipmentExportError(
                f"Indus Part is missing for HU {hu}."
            )

        prefix = _part_prefix(part)
        description = description_map.get(prefix)

        if not description:
            raise ShipmentExportError(
                "The shipment template has no Description of Goods "
                f"mapping for part {part} / HU {hu}."
            )

        seen_hus.add(hu)

        normalised.append(
            {
                "handling_unit": hu,
                "description": description,
                "customer_part": customer,
                "indus_part": part,
                "pallet_qty": quantity,
            }
        )

    return normalised


def build_shipment_export(loading):
    if not TEMPLATE_PATH.is_file():
        raise ShipmentExportError(
            "The advance shipment Excel template is missing."
        )

    if not isinstance(loading, dict):
        raise ShipmentExportError(
            "Loading export data is invalid."
        )

    reference_id = _text(loading.get("reference_id"))

    if not reference_id:
        raise ShipmentExportError(
            "The loading Reference ID is missing."
        )

    status = _text(loading.get("status")).upper()

    if status != "COMPLETED":
        raise ShipmentExportError(
            "Only completed loadings can be exported."
        )

    completed_date = _loading_date(
        loading.get("completed_at")
    )

    wb = load_workbook(
        TEMPLATE_PATH,
        data_only=False,
        keep_links=True,
    )

    if "Shipment" not in wb.sheetnames:
        raise ShipmentExportError(
            "The template does not contain the Shipment sheet."
        )

    ws = wb["Shipment"]
    description_map = _template_description_map(ws)
    items = _normalise_items(
        loading.get("items"),
        description_map,
    )

    extra_rows = max(
        0,
        len(items) - (
            BASE_DATA_END_ROW - DATA_START_ROW + 1
        ),
    )

    _insert_shipment_rows(ws, extra_rows)

    data_end_row = max(
        BASE_DATA_END_ROW,
        DATA_START_ROW + len(items) - 1,
    )
    total_row = BASE_TOTAL_ROW + extra_rows
    signature_header_row = 40 + extra_rows
    signature_end_row = 45 + extra_rows

    ws["F2"] = reference_id

    ws["J6"] = completed_date
    ws["J6"].number_format = "d-mmm-yy"

    ws["J7"] = None
    ws["J9"] = None

    customers = []
    for item in items:
        customer = item["customer_part"]
        if customer and customer not in customers:
            customers.append(customer)

    ws["J8"] = ", ".join(customers)
    ws["J8"].number_format = "General"

    for row in range(DATA_START_ROW, data_end_row + 1):
        for column in range(1, 8):
            ws.cell(row, column).value = None

    for sequence, item in enumerate(items, start=1):
        row = DATA_START_ROW + sequence - 1

        ws.cell(row, 1).value = sequence
        ws.cell(row, 2).value = item["description"]
        ws.cell(row, 3).value = item["customer_part"]
        ws.cell(row, 4).value = item["indus_part"]

        # One workbook row represents one HU/pallet.
        ws.cell(row, 5).value = item["pallet_qty"]
        ws.cell(row, 6).value = 1
        ws.cell(row, 7).value = item["pallet_qty"]

    ws.cell(total_row, 5).value = "Total"
    ws.cell(total_row, 6).value = (
        f"=SUM(F{DATA_START_ROW}:F{data_end_row})"
    )
    ws.cell(total_row, 7).value = (
        f"=SUM(G{DATA_START_ROW}:G{data_end_row})"
    )

    # PO number and invoice number must remain empty.
    ws["J7"] = None
    ws["J9"] = None

    # Keep Name and Sign headings, clear only their value cells.
    for row in range(
        signature_header_row + 1,
        signature_end_row + 1,
    ):
        ws.cell(row, 4).value = None
        ws.cell(row, 5).value = None

    # Remove the broken, invisible #REF formulas outside the print area.
    for row in range(11, data_end_row + 1):
        for column in range(15, 29):
            ws.cell(row, column).value = None

    ws.print_area = (
        f"A1:N{BASE_PRINT_END_ROW + extra_rows}"
    )

    calculation = getattr(wb, "calculation", None)
    if calculation is not None:
        calculation.fullCalcOnLoad = True
        calculation.forceFullCalc = True

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    safe_reference = secure_filename(reference_id)
    if not safe_reference:
        safe_reference = "loading"

    filename = (
        f"Advance-Shipment-{safe_reference}-"
        f"{completed_date.isoformat()}.xlsx"
    )

    return output, filename
